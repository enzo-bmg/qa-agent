#!/usr/bin/env python3
"""
extract_docs.py — Converte PDF e PPTX para Markdown.

Uso no laptop BMG (só terminal + VS Code):
    # instala as deps num venv isolado (não precisa admin)
    uv venv .venv-extract
    # Windows:  .venv-extract\\Scripts\\activate
    # Linux/Mac: source .venv-extract/bin/activate
    uv pip install pymupdf4llm python-pptx

    # roda apontando pra pasta dos docs do Orange Box
    python extract_docs.py "C:\\caminho\\para\\OrangeBox"
    # ou saída customizada:
    python extract_docs.py "C:\\OrangeBox" --out "C:\\OrangeBox\\md"

O que faz:
  - varre a pasta (recursivo) atrás de .pdf, .ppt, .pptx
  - PDF  -> markdown via pymupdf4llm (preserva headings/tabelas)
  - PPTX -> markdown (título + bullets + texto de tabelas por slide)
  - salva um .md por arquivo, mantendo o nome original
  - gera um _INDEX.md com links pra todos os .md extraídos

Não envia nada pra fora — processamento 100% local.
"""

import argparse
import sys
from pathlib import Path


# Perguntas do checklist (20260916_checklist_docs_orangebox.md), agrupadas por
# tipo de doc. A chave é um trecho do nome do arquivo (lowercase) pra casar.
CHECKLIST = {
    "normativa": {
        "titulo": "Normativa de mudança → schema + mock da GMUD",
        "perguntas": [
            "Nome real do identificador da GMUD (gmud_id? change_number? CHG...?)",
            "Campos obrigatórios de uma GMUD (listar todos)",
            "Tipos de mudança (enum real: normal / emergencial / padrão / pré-aprovada?)",
            "Campo de sistemas afetados — como se chama? é CI do CMDB? texto livre? lista?",
            "Campo de janela de execução (início/fim, formato de data)",
            "Campo de responsável / solicitante / aprovador",
            "Campo de ambiente (produção / homologação / DR)?",
            "Status possíveis da GMUD (enum)",
            "Existe campo linkando ao PR/repositório (Azure DevOps)?",
            "Existe campo de score/risco/criticidade já hoje?",
        ],
    },
    "rollback": {
        "titulo": "Rollback de mudança → schema + mock de rollback",
        "perguntas": [
            "Existe registro estruturado de rollback? Onde fica (ServiceNow / Azure / planilha)?",
            "Campos do registro de rollback (causa raiz, impacto, tempo, sistemas)",
            "Categoria de falha é padronizada (enum)? Quais valores?",
            "Registra tempo de retrabalho / indisponibilidade?",
            "Vincula rollback à GMUD original (campo de referência)?",
            "Tem lições aprendidas / post-mortem estruturado?",
        ],
    },
    "homolog": {
        "titulo": "Teste de homologação → FORMATO DE SAÍDA do agente",
        "perguntas": [
            "Formato/template real que o QA usa pra registrar cenário de teste",
            "Campos de um caso de teste (ID, título, pré-condição, passos, resultado esperado, evidência?)",
            "Como registram critério de aprovação/reprovação",
            "Onde registram evidências (screenshot, log)?",
            "Tem classificação de prioridade/severidade dos testes? Quais valores?",
            "Existe checklist obrigatório de homologação por tipo de sistema?",
        ],
    },
    "treina": {
        "titulo": "Treinamento gestão de mudança → prompt + workflow",
        "perguntas": [
            "Vocabulário/termos internos do BMG (usar no prompt pra soar nativo)",
            "Passo a passo que o QA segue hoje (do recebimento à aprovação)",
            "Gargalos mencionados no processo",
            "Quem aprova o quê (níveis de aprovação por risco)",
            "Regras de compliance/segregação de função relevantes",
        ],
    },
}


def match_checklist(nome_arquivo: str):
    """Retorna o bloco de checklist que casa com o nome do arquivo, ou None."""
    low = nome_arquivo.lower()
    for chave, bloco in CHECKLIST.items():
        if chave in low:
            return bloco
    return None


def build_prompt(nome_arquivo: str, conteudo_md: str) -> str:
    """Monta um prompt pronto pra colar no Flow, cruzando doc + perguntas."""
    bloco = match_checklist(nome_arquivo)
    if bloco:
        perguntas = "\n".join(f"{i}. {q}" for i, q in enumerate(bloco["perguntas"], 1))
        foco = bloco["titulo"]
    else:
        # doc não reconhecido — pede um resumo estruturado genérico
        foco = "Documento normativo (tipo não identificado pelo nome)"
        perguntas = (
            "1. Quais campos/estruturas de dados este documento define?\n"
            "2. Quais enums, status ou categorias padronizadas aparecem?\n"
            "3. Que termos/vocabulário interno vale reaproveitar?\n"
            "4. Qual o passo a passo/processo descrito?"
        )

    return f"""Você é um analista ajudando a modelar um agente de QA para o Banco BMG.
Abaixo está o conteúdo extraído de um documento normativo interno ({foco}).

TAREFA: leia o documento e responda objetivamente cada pergunta abaixo.
- Se a informação existir no texto, cite o valor exato (nome do campo, enum, etc.).
- Se NÃO existir, responda "não consta no documento".
- Não invente. Priorize precisão sobre completude.

=== PERGUNTAS ===
{perguntas}

=== DOCUMENTO: {nome_arquivo} ===
{conteudo_md}
"""


def pdf_to_md(pdf_path: Path) -> str:
    """PDF -> markdown. Tenta pymupdf4llm; fallback pra pymupdf puro."""
    try:
        import pymupdf4llm
        return pymupdf4llm.to_markdown(str(pdf_path))
    except ImportError:
        pass

    # fallback: pymupdf (fitz) extraindo texto puro por página
    try:
        import fitz  # pymupdf
    except ImportError:
        return "> [ERRO] Instale `pymupdf4llm` ou `pymupdf` para ler PDF.\n"

    doc = fitz.open(str(pdf_path))
    partes = []
    for i, page in enumerate(doc, start=1):
        partes.append(f"\n\n---\n\n## Página {i}\n\n{page.get_text()}")
    doc.close()
    return "".join(partes)


def pptx_to_md(pptx_path: Path) -> str:
    """PPTX -> markdown: um bloco por slide (título, bullets, tabelas)."""
    try:
        from pptx import Presentation
    except ImportError:
        return "> [ERRO] Instale `python-pptx` para ler PPTX.\n"

    prs = Presentation(str(pptx_path))
    linhas = []

    for idx, slide in enumerate(prs.slides, start=1):
        linhas.append(f"\n\n---\n\n## Slide {idx}\n")

        # título (se houver placeholder de título)
        titulo = None
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                titulo = slide.shapes.title.text.strip()
        except Exception:
            titulo = None
        if titulo:
            linhas.append(f"\n### {titulo}\n")

        for shape in slide.shapes:
            # pula o título (já tratado)
            if titulo and shape == slide.shapes.title:
                continue

            # texto com bullets
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = "".join(run.text for run in para.runs).strip()
                    if not txt:
                        continue
                    indent = "  " * getattr(para, "level", 0)
                    linhas.append(f"{indent}* {txt}")

            # tabelas
            if shape.has_table:
                tbl = shape.table
                rows = list(tbl.rows)
                if rows:
                    # header
                    header = [c.text.strip() for c in rows[0].cells]
                    linhas.append("\n| " + " | ".join(header) + " |")
                    linhas.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for r in rows[1:]:
                        cells = [c.text.strip().replace("\n", " ") for c in r.cells]
                        linhas.append("| " + " | ".join(cells) + " |")
                    linhas.append("")

            # notas do slide (às vezes tem info útil)
        try:
            if slide.has_notes_slide:
                notas = slide.notes_slide.notes_text_frame.text.strip()
                if notas:
                    linhas.append(f"\n> **Notas do slide:** {notas}\n")
        except Exception:
            pass

    return "\n".join(linhas)


def main():
    ap = argparse.ArgumentParser(description="Extrai PDF/PPTX para Markdown (local).")
    ap.add_argument("pasta", help="Pasta com os documentos (busca recursiva)")
    ap.add_argument("--out", help="Pasta de saída (default: <pasta>/_md_extraido)")
    ap.add_argument(
        "--prompt",
        action="store_true",
        help="Além do .md, gera um _prompt.txt por doc (doc + perguntas do checklist) pronto pra colar no Flow",
    )
    args = ap.parse_args()

    base = Path(args.pasta).expanduser().resolve()
    if not base.exists():
        print(f"[ERRO] Pasta não existe: {base}")
        sys.exit(1)

    out_dir = Path(args.out).expanduser().resolve() if args.out else base / "_md_extraido"
    out_dir.mkdir(parents=True, exist_ok=True)

    arquivos = sorted(
        p for p in base.rglob("*")
        if p.suffix.lower() in {".pdf", ".ppt", ".pptx"} and out_dir not in p.parents
    )

    if not arquivos:
        print(f"[AVISO] Nenhum PDF/PPT/PPTX encontrado em {base}")
        sys.exit(0)

    print(f"Encontrados {len(arquivos)} arquivo(s). Saída em: {out_dir}\n")

    index = ["# Índice — Documentos extraídos\n"]

    for arq in arquivos:
        print(f"  -> {arq.name} ...", end=" ", flush=True)
        try:
            if arq.suffix.lower() == ".pdf":
                md = pdf_to_md(arq)
            elif arq.suffix.lower() == ".pptx":
                md = pptx_to_md(arq)
            else:  # .ppt (formato antigo) — python-pptx não lê
                md = (
                    f"> [PULADO] `.ppt` (formato antigo) não é suportado pelo python-pptx.\n"
                    f"> Abra `{arq.name}` no PowerPoint/LibreOffice e salve como `.pptx`, depois rode de novo.\n"
                )
                print("PULADO (.ppt antigo)")
                # ainda salva o aviso
        except Exception as e:
            md = f"> [ERRO ao processar] {arq.name}: {e}\n"
            print(f"ERRO: {e}")
        else:
            if arq.suffix.lower() != ".ppt":
                print("ok")

        destino = out_dir / (arq.stem + ".md")
        destino.write_text(f"# {arq.name}\n\n{md}", encoding="utf-8")
        rel = destino.relative_to(out_dir)
        index.append(f"* [{arq.name}](./{rel.as_posix()})")

        if args.prompt and arq.suffix.lower() != ".ppt":
            prompt_txt = build_prompt(arq.name, md)
            (out_dir / (arq.stem + "_prompt.txt")).write_text(prompt_txt, encoding="utf-8")

    (out_dir / "_INDEX.md").write_text("\n".join(index) + "\n", encoding="utf-8")
    print(f"\nPronto. Abra {out_dir / '_INDEX.md'} no VS Code.")


if __name__ == "__main__":
    main()
